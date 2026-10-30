#!/usr/bin/env python3
"""Perbarui data tambahan kodeBMN.html. Jalankan:  python3 update.py

1) Foto       : isi folder www/foto/  (nama file = kode 10 digit, mis. 3050101001.jpg)
                -> membuat ulang www/foto/daftar.js
2) Keterangan : edit keterangan.xlsx (sheet "Keterangan": kode, nama, alias, keterangan)
                -> membuat ulang www/keterangan.js  (butuh: pip install openpyxl)
Skrip ini juga dijalankan otomatis oleh GitHub saat membangun APK.
"""
import os, re, sys, json
B = os.path.dirname(os.path.abspath(__file__))

# ---- Foto ----
W = os.path.join(B, 'www')
D = os.path.join(W, 'foto')
ok = {'.jpg', '.jpeg', '.png', '.webp'}
man, skip = {}, []
for f in sorted(os.listdir(D)):
    n, e = os.path.splitext(f)
    if e.lower() not in ok:
        continue
    if re.fullmatch(r'\d{10}', n):
        man[n] = f'foto/{f}'
    else:
        skip.append(f)
with open(os.path.join(D, 'daftar.js'), 'w', encoding='utf8') as f:
    f.write('// Dibuat otomatis oleh update.py — jangan diedit manual.\n'
            'const PHOTOS = ' + json.dumps(man, indent=1, sort_keys=True) + ';\n')
print(f'Foto: {len(man)} terdaftar.')
for s in skip:
    print('  dilewati (nama bukan 10 digit kode):', s)

# ---- Keterangan & alias ----
try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit('Perlu openpyxl. Jalankan dulu:  pip install openpyxl')
kamus, seen, warn = {}, set(), []
wb = load_workbook(os.path.join(B, 'keterangan.xlsx'), read_only=True, data_only=True)
ws = wb['Keterangan'] if 'Keterangan' in wb.sheetnames else wb.worksheets[0]
head = None
for no, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if head is None:
        head = [str(c or '').strip().lower() for c in row]
        for need in ('kode', 'alias', 'keterangan'):
            if need not in head:
                sys.exit(f'Kolom "{need}" tidak ditemukan di baris judul keterangan.xlsx')
        continue
    r = dict(zip(head, row))
    def cell(k):
        v = r.get(k)
        if v is None:
            return ''
        if isinstance(v, float) and v.is_integer():
            v = int(v)          # Excel kadang menyimpan kode sebagai angka (3010304003.0)
        return str(v).strip()
    kode, alias, ket = cell('kode'), cell('alias'), cell('keterangan')
    if not kode and not alias and not ket:
        continue
    if not re.fullmatch(r'\d{10}', kode):
        warn.append(f'baris {no}: kode "{kode}" bukan 10 digit, dilewati'); continue
    if kode in seen:
        warn.append(f'baris {no}: kode {kode} ganda, yang terakhir dipakai')
    seen.add(kode)
    if alias or ket:
        kamus[kode] = [alias, ket]
    else:
        kamus.pop(kode, None)
wb.close()
with open(os.path.join(W, 'keterangan.js'), 'w', encoding='utf8') as f:
    f.write('// Dibuat otomatis oleh update.py dari keterangan.xlsx — jangan diedit manual.\n'
            '// KETERANGAN[kode] = [alias, keterangan]\n'
            'const KETERANGAN = ' + json.dumps(kamus, ensure_ascii=False, indent=1, sort_keys=True) + ';\n')
print(f'Keterangan: {len(kamus)} kode ({sum(1 for v in kamus.values() if v[0])} punya alias, {sum(1 for v in kamus.values() if v[1])} punya keterangan).')
for w in warn:
    print('  peringatan:', w)

# ---- Cek kode ada di data utama ----
m = re.search(r'const DATA = (\[.*?\]);\n', open(os.path.join(W, 'index.html'), encoding='utf8').read(), re.S)
if m:
    semua = {r[0] for r in json.loads(m.group(1))}
    for k in sorted(set(man) | set(kamus)):
        if k not in semua:
            print('  peringatan: kode', k, 'tidak ada di data utama')
